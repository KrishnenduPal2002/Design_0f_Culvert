"""Inventory workbook formulas and references without executing them."""
import argparse
import csv
from pathlib import Path
import openpyxl


def audit_workbook(workbook_path: str, output_csv: str) -> dict:
    workbook = openpyxl.load_workbook(workbook_path, data_only=False, read_only=True, keep_links=True)
    totals = {"sheets": len(workbook.sheetnames), "formulas": 0,
              "broken_ref_formulas": 0, "external_ref_formulas": 0}
    out = Path(output_csv)
    out.parent.mkdir(parents=True, exist_ok=True)
    with out.open("w", newline="", encoding="utf-8-sig") as file:
        writer = csv.writer(file)
        writer.writerow(["sheet", "cell", "formula", "contains_ref_error", "contains_external_link"])
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    formula = cell.value
                    if cell.data_type != "f" or not isinstance(formula, str):
                        continue
                    broken = "#REF!" in formula
                    external = "[" in formula and "]" in formula
                    totals["formulas"] += 1
                    totals["broken_ref_formulas"] += int(broken)
                    totals["external_ref_formulas"] += int(external)
                    writer.writerow([sheet.title, cell.coordinate, formula, broken, external])
    workbook.close()
    return totals


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", nargs="?", default="data/original_design.xlsx")
    parser.add_argument("--output", default="outputs/formula_inventory.csv")
    args = parser.parse_args()
    print(audit_workbook(args.workbook, args.output))
    print(f"Formula inventory saved to {args.output}")
